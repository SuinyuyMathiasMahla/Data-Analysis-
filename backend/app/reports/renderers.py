"""Report renderers: PDF (ReportLab), DOCX (python-docx), XLSX (openpyxl), CSV.

All four render the same ``content`` dictionary (built by
``app.services.report_service.build_report_content`` and stored as the report
version's snapshot), so every format tells the same story from the same
numbers. Renderers contain no statistics and no database access.

``images`` maps a figure's ``image_key`` to PNG bytes.
"""
from __future__ import annotations

import csv
import re
import io
import math
from datetime import datetime, timezone
from typing import Any

SECTION_TITLES = [
    ("executive_summary", "1. Abstract"),
    ("background", "2. Background and literature review"),
    ("questions", "3. Research questions and hypotheses"),
    ("methods", "4. Methodology"),
    ("cleaning", "5. Data cleaning summary"),
    ("descriptive", "6. Results: descriptive statistics"),
    ("inferential", "7. Results: statistical tests"),
    ("review_findings", "Review of findings"),
    ("discussion", "8. Discussion"),
    ("limitations", "9. Limitations"),
    ("conclusions", "10. Conclusions"),
    ("recommendations", "11. Recommendations"),
    ("references", "12. References"),
    ("appendices", "Appendices"),
]
INTEGRITY_NOTE = ("The numbers in this report were calculated by the platform from the dataset named on the cover page. "
                  "The written explanations were prepared from those numbers. The researcher checks them and is responsible for the conclusions.")
# What each section is for, in plain words, shown at the start of the section.
INTROS = {
    "executive_summary": "This short summary gives the answer to each research question. Each answer is explained in full later in the report.",
    "background": "This section explains the problem the study addresses, the aim of the study, and what earlier research has found about the topic.",
    "questions": "This section lists what the study set out to find out. A research question is the question the study answers. "
                 "A hypothesis is a statement that can be tested with data: the null hypothesis (H0) says there is no effect or no difference, "
                 "and the alternative hypothesis (H1) says there is one.",
    "methods": "This section explains how the study was done: what was studied, where the data came from, what was measured, "
               "and which statistical tests were used and why.",
    "cleaning": "Raw data often contain mistakes such as repeated rows, impossible values and gaps. This section lists every change made "
                "to the data before the analysis, why it was made, and how many records it affected.",
    "descriptive": "This section describes the data before any hypothesis is tested: averages, spreads and counts for each variable. "
                   "It shows what the sample looks like.",
    "inferential": "This section reports the statistical tests. For each test it gives the result, how large the effect is, the range the true "
                   "effect is likely to fall in, and whether the conditions the test needs were met.",
    "review_findings": "This section reviews each finding in turn: the question it answers, the variables, the test and why it was chosen, "
                       "how the calculation was done, what the numbers mean, what they suggest, and the chart for the finding.",
    "discussion": "This section explains what each finding means in plain terms.",
    "limitations": "No study is perfect. This section lists the weaknesses that affect how far the results can be trusted.",
    "conclusions": "This section states what the study has shown, one conclusion per research question, each tied to the statistic behind it.",
    "recommendations": "This section gives practical actions based on the findings. Each action says what should be done, who should do it, and when.",
    "references": "This section lists the sources cited in the report and the statistical methods used.",
    "appendices": "The appendices hold the detailed tables behind the report, a glossary of the terms used, and the analysis settings, "
                  "so that the results can be checked and repeated.",
}
GLOSSARY = [
    ("Research question", "The question the study sets out to answer."),
    ("Hypothesis", "A statement that can be tested with data. The null hypothesis (H0) says there is no effect or difference; the alternative (H1) says there is one."),
    ("Variable", "Something measured for each person or item in the data, such as a score, an age or a category."),
    ("Test statistic", "The number a statistical test calculates from the data (for example t, F, chi-square or r). Larger values usually mean a stronger pattern."),
    ("Degrees of freedom (df)", "A number based on the sample size that the test needs to judge how much information the data hold."),
    ("p-value", "The chance of getting a result at least as extreme as the one observed if there were really no effect. "
                "In this report a p-value below 0.05 is treated as evidence against the null hypothesis."),
    ("Significance level (alpha)", "The cut-off for the p-value, set before the analysis. This report uses the value given in each hypothesis."),
    ("Effect size", "How large a difference or relationship is, separate from whether it is statistically significant. A small p-value with a small effect size means the effect is real but modest."),
    ("95% confidence interval (CI)", "A range of values for the true effect. If the study were repeated many times, about 95 in 100 of the ranges would contain the true value."),
    ("Correlation", "How closely two measurements move together, from -1 (opposite directions) through 0 (no link) to +1 (same direction)."),
    ("Assumption", "A condition the data should meet for a test to give reliable results, for example that values are roughly normally spread."),
    ("Non-parametric test", "A test that makes fewer assumptions about the shape of the data. It is used when the usual test's assumptions are not met."),
    ("Holm adjustment", "A correction applied when several tests are run together, so that chance results are not accepted too easily."),
    ("Missing value", "A cell in the data with no recorded value."),
    ("Outlier", "A value far from most of the others. It may be a real extreme case or a recording mistake, so it is checked rather than removed automatically."),
]

def fmt(value: Any, digits: int = 3) -> str:
    """Display formatting shared by all renderers."""
    if value is None:
        return "-"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return "-"
        if value.is_integer() and abs(value) < 1e9:
            return str(int(value))
        if value != 0 and abs(value) < 0.001:
            return "< 0.001" if value > 0 else f"{value:.2e}"
        return f"{value:.{digits}f}"
    return str(value)


def _sections(content: dict) -> list[tuple[str, str]]:
    wanted = content.get("sections")
    return [(k, t) for k, t in SECTION_TITLES if not wanted or k in wanted]


def _cover_rows(content: dict) -> list[tuple[str, str]]:
    """The title page: author, institution, date, degree programme and supervisor."""
    p = content.get("project", {})
    rows = [(p.get("researcher_label") or "Researcher", p.get("researcher")), ("Institution", p.get("institution")),
            ("Department", p.get("department")),
            ("Degree programme", {"undergraduate": "Undergraduate", "masters": "Master's", "phd": "PhD"}.get(p.get("academic_level"), p.get("academic_level"))),
            ("Supervisor", ", ".join(p.get("supervisors") or []) or None),
            ("Research field", p.get("research_field")), ("Date", content.get("generated_at", "")[:10])]
    return [(k, str(v)) for k, v in rows if v]


def _paragraphs(text: str | None) -> list[str]:
    return [line.strip() for line in (text or "").split("\n") if line.strip()]


def _inferential_blocks(items: list[dict], add, prefix: str) -> None:
    """Each statistical test: heading, what was compared, the result, effect sizes, estimates, assumptions, warnings."""
    for i, item in enumerate(items, 1):
        tag = " / ".join(x for x in (item.get("rq_code"), item.get("hypothesis_code")) if x)
        add(("h2", f"{prefix}{i} {item['method_name']}" + (f" ({tag})" if tag else "")))
        add(("p", f"Variables compared: {item['variables_text']}. Number of records used: n = {item.get('n')}."))
        if item.get("apa"):
            add(("p", f"Result: {item['apa']}."))
        if item.get("summary"):
            add(("p", item["summary"]))
        if item.get("effect_sizes"):
            add(("table", {"keep": True, "title": "Size of the effect", "columns": ["Measure", "Value", "95% CI low", "95% CI high", "Size"],
                           "rows": [[e["name"], e.get("value"), e.get("ci_low"), e.get("ci_high"), e.get("magnitude") or ""] for e in item["effect_sizes"]]}))
        if item.get("estimates"):
            add(("table", {"keep": True, "title": "Estimates with 95% confidence intervals", "columns": ["Estimate", "Value", "95% CI low", "95% CI high"],
                           "rows": [[e["name"], e.get("value"), e.get("ci_low"), e.get("ci_high")] for e in item["estimates"]]}))
        if item.get("assumptions"):
            add(("table", {"keep": True, "title": "Conditions the test needs (assumption checks)", "columns": ["Condition", "Check used", "Statistic", "p", "Met?", "Note"],
                           "rows": [[a["name"], a.get("test") or "", a.get("statistic"), a.get("p_value"),
                                     "Yes" if a.get("satisfied") is True else "No" if a.get("satisfied") is False else "Judge", a.get("note") or ""]
                                    for a in item["assumptions"]]}))
        for t in item.get("tables", []):
            add(("table", t))
        for w in item.get("warnings", []):
            add(("note", f"Warning: {w}"))


def _tests_in(content: dict, categories: tuple[str, ...]) -> list[dict]:
    return [i for i in content.get("inferential") or [] if i.get("category") in categories]


BIVARIATE = ("Correlation", "Group comparison", "Association")
MULTIVARIATE = ("Regression", "Multivariate", "Advanced")
MODEL_CATEGORIES = ("Regression", "Multivariate", "Advanced")


def _points(t: dict) -> list[tuple[str, Any]]:
    """A small table in point form: one bullet per row, 'row label: column value; ...'. Keeps the figures, drops the grid."""
    out: list[tuple[str, Any]] = []
    if t.get("title"):
        out.append(("p", f"{t['title']}:"))
    cols = t.get("columns") or []
    for row in t.get("rows") or []:
        if not row:
            continue
        details = "; ".join(f"{c}: {v}" for c, v in zip(cols[1:], row[1:]) if v not in (None, ""))
        out.append(("bullet", f"{row[0]}: {details}" if details else str(row[0])))
    if t.get("note"):
        out.append(("note", t["note"]))
    return out


def _text_blocks(content: dict, key: str, raw_tables: bool = False) -> list[tuple[str, Any]]:
    """The blocks of one report section (see SECTION_TITLES), as ('h2' | 'p' | 'bullet' | 'table' | 'figure' | 'note', payload).
    Only the tables marked keep are drawn as tables; the rest are written as points. raw_tables keeps every table (for the workbook)."""
    blocks = _section_blocks(content, key)
    if raw_tables:
        return blocks
    out: list[tuple[str, Any]] = []
    for kind, payload in blocks:
        if kind == "table" and not payload.get("keep"):
            out += _points(payload)
        else:
            out.append((kind, payload))
    return out


def _review_lines(add, lines: list[str], chart: dict | None) -> None:
    """A finding's review: headings for the eight parts, bullets as bullets, and its chart at part 8."""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith("FINDING "):
            add(("h2", line))
        elif re.match(r"^\d\. [A-Z ,&/-]+:$", line):
            add(("h2", line.rstrip(":")))
        elif line.startswith("- "):
            add(("bullet", line[2:]))
        else:
            add(("p", line))
        if line.startswith("8. VISUALISATION") and chart:
            add(("figure", {"image_key": chart["image_key"], "title": f"Chart for this finding: {chart['title']}",
                            "note": chart.get("note")}))


def _layout_lines(add, text: str, figure: dict | None = None) -> None:
    """Write a conclusion or recommendation as the lines it was written in: one sentence per line, bullets kept as bullets.
    ``figure`` (the chart that explains the conclusion) is placed just before the conclusive statement, after what was found."""
    for line in (text or "").split("\n"):
        line = line.strip()
        if not line:
            continue
        if figure is not None and line.startswith("Conclusion on the Hypotheses:"):
            add(("figure", figure))
            figure = None
        if line.startswith("- "):
            add(("bullet", line[2:]))
        else:
            add(("p", line))
    if figure is not None:
        add(("figure", figure))


def _section_blocks(content: dict, key: str) -> list[tuple[str, Any]]:
    blocks: list[tuple[str, Any]] = []
    add = blocks.append
    project = content.get("project") or {}
    dataset = content.get("dataset") or {}
    quality = content.get("data_quality") or {}
    summary = quality.get("summary") or {}
    variables = dataset.get("variables") or []
    tests = content.get("inferential") or []
    hyps = content.get("hypotheses") or []
    figures = content.get("figures") or []
    cleaning = quality.get("cleaning") or []
    rows = dataset.get("rows")
    details = content.get("details") or {}
    alpha = hyps[0]["alpha"] if hyps else 0.05
    if key in INTROS:
        add(("p", INTROS[key]))

    if key == "executive_summary":
        for sentence in content.get("executive_summary") or []:
            add(("p", sentence))
        # the same verdict table as the one-page summary, so page 1 of the report answers the questions on its own
        if hyps:
            add(("table", {"keep": True, "title": "Answers at a glance", "columns": ["Hypothesis", "Test", "Statistic", "p", "Adjusted p (Holm)", "Decision"],
                           "rows": [[h["code"] + (" (primary)" if h.get("primary") else ""), h.get("test") or "Not tested", h.get("statistic_text") or "-",
                                     h.get("p_value"), h.get("p_adjusted"), h["decision_text"]] for h in hyps]}))
        else:
            add(("p", "No hypotheses have been defined, so there are no answers to summarise yet."))

    elif key == "background":
        add(("h2", "2.1 The problem"))
        add(("p", project.get("research_problem") or "The research problem has not been recorded in the project."))
        add(("h2", "2.2 Field of study"))
        add(("p", f"The study belongs to the field of {project['research_field']}." if project.get("research_field")
            else "The field of study has not been recorded in the project."))
        add(("h2", "2.3 Aim of the study"))
        add(("p", project.get("general_objective") or "The aim has not been recorded in the project."))
        add(("h2", "2.4 What earlier research says"))
        background = _paragraphs(details.get("background"))
        if background:
            for paragraph in background:
                add(("p", paragraph))
        else:
            add(("note", "The literature review has not been entered yet. Add it under Report details on the Report page, and it will appear here."))

    elif key == "questions":
        add(("h2", "3.1 Objectives"))
        objectives = project.get("objectives") or []
        for text in objectives:
            add(("bullet", text))
        if not objectives:
            add(("p", "No objectives have been recorded."))
        add(("h2", "3.2 Research questions"))
        rqs = content.get("research_questions") or []
        if rqs:
            add(("table", {"keep": True, "title": "Research questions", "columns": ["Code", "Research question", "Status"],
                           "rows": [[rq["code"], rq["text"], rq["status_text"]] for rq in rqs]}))
        else:
            add(("p", "No research questions have been recorded."))
        add(("h2", "3.3 Hypotheses"))
        if hyps:
            add(("p", f"Each hypothesis is judged at a significance level (alpha) of {alpha:g}. "
                      "If the p-value is below this level, the null hypothesis is rejected."))
            add(("table", {"keep": True, "title": "Hypotheses and planned tests", "columns": ["Hypothesis", "Null hypothesis (H0)", "Alternative hypothesis (H1)", "Planned test"],
                           "rows": [[h["code"], h["null_text"], h["alternative_text"], h.get("test") or "Not yet run"] for h in hyps]}))
        else:
            add(("p", "No hypotheses have been defined."))

    elif key == "methods":
        add(("h2", "4.1 Study design and participants"))
        design = [f"Study design: {project['study_design']}." if project.get("study_design") else None,
                  f"Study population: {project['study_population']}." if project.get("study_population") else None,
                  f"Sampling technique: {project['sampling_technique']}." if project.get("sampling_technique") else None,
                  f"Planned sample size: {project['sample_size']}." if project.get("sample_size") else None]
        lines = [x for x in design if x]
        for line in lines:
            add(("bullet", line))
        if dataset:
            add(("p", f"The analysis used the dataset '{dataset.get('name')}', which has {rows} records and {dataset.get('columns')} variables after cleaning."))
        if not lines and not dataset:
            add(("p", "The study design has not been recorded in the project."))
        add(("h2", "4.2 Data collection"))
        add(("p", project.get("data_collection_method") or "The data collection method has not been recorded in the project."))
        add(("h2", "4.3 Variables"))
        add(("p", "A variable is a characteristic that was measured. The role says how the variable was used: the dependent variable is the outcome, "
                  "and the independent variable is the factor thought to affect it."))
        if variables:
            add(("table", {"keep": True, "title": "Variables and their roles", "columns": ["Variable", "Label", "Type", "Role"],
                           "rows": [[v["name"], v.get("label") or "", v["data_type"], v.get("role") or "other"] for v in variables]}))
        else:
            add(("p", "No variables are recorded."))
        add(("h2", "4.4 Statistical tests and why they were chosen"))
        add(("p", "Each test was chosen from the research question, the type of each variable, and whether the data met the conditions the test needs. "
                  "Where a condition was not met, a test that does not need it is reported beside the usual test."))
        for line in [
            "Cleaning: repeated rows, impossible values and missing-value codes were handled before any test (Section 5).",
            "Descriptive results: summary figures for each variable (Section 6).",
            "Conditions: the assumptions of each test were checked, for example that values are roughly normally spread (Appendix A).",
            "Effect sizes and precision: effect sizes and 95% confidence intervals are reported with each test where the method provides them.",
            f"Decisions: each hypothesis is judged at a significance level of {alpha:g}.",
            "Several tests: where more than one test is reported, p-values are adjusted with the Holm method (Section 7).",
        ]:
            add(("bullet", line))
        software = (content.get("methodology") or {}).get("software") or {}
        if software:
            add(("bullet", "Software used: " + ", ".join(f"{k} {v}" for k, v in software.items()) + "."))
        add(("h2", "4.5 Ethics"))
        if details.get("ethics_committee") or details.get("ethics_reference"):
            add(("bullet", f"Ethics committee: {details.get('ethics_committee') or 'not recorded'}."))
            add(("bullet", f"Approval reference and date: {details.get('ethics_reference') or 'not recorded'}."))
            add(("bullet", f"Consent: {details.get('consent_method') or 'not recorded'}."))
        else:
            add(("note", "The ethics approval has not been entered yet. Add the committee, approval reference and consent method under Report details on the Report page."))

    elif key == "cleaning":
        add(("h2", "5.1 Cleaning steps"))
        if cleaning:
            add(("table", {"keep": True, "title": "Cleaning steps",
                           "columns": ["#", "Step", "Why it was done", "Records before", "Records after", "Cells changed"],
                           "rows": [[i + 1, op["label"], op["reason"], op["rows_before"], op["rows_after"], op["cells_changed"]]
                                    for i, op in enumerate(cleaning)]}))
            add(("p", f"{cleaning[-1]['rows_after']} records remained after cleaning, from {cleaning[0]['rows_before']} records received."))
        else:
            add(("p", "No cleaning steps were applied, so the data were analysed as they were received."))
        add(("h2", "5.2 Missing values"))
        if summary:
            add(("p", f"{summary.get('missing_cells', 0)} cells are missing, which is {fmt(summary.get('missing_pct', 0), 2)}% of all cells."))
        else:
            add(("p", "No data-quality check has been run."))
        missing = [v for v in variables if v.get("n_missing")]
        if missing:
            add(("table", {"keep": True, "title": "Missing values after cleaning", "columns": ["Variable", "Missing (n)", "Missing (%)"],
                           "rows": [[v["name"], v["n_missing"], round(100 * v["n_missing"] / rows, 1) if rows else None] for v in missing]}))
        add(("h2", "5.3 Repeated rows, outliers and invalid values"))
        if summary:
            add(("p", f"{summary.get('duplicate_rows', 0)} repeated records were found. {summary.get('possible_outliers', 0)} values were flagged as possible outliers "
                      f"(values far from most others) and {summary.get('invalid_cells', 0)} values were flagged as invalid or impossible."))
            for issue in [i for i in quality.get("issues") or [] if "outlier" in str(i.get("message", "")).lower()][:10]:
                add(("bullet", issue["message"]))

    elif key == "descriptive":
        add(("h2", "6.1 Summary figures"))
        descriptive = content.get("descriptive") or []
        tables = [t for item in descriptive for t in item.get("tables", [])]
        if not tables:
            add(("p", "No descriptive analysis has been run."))
        for t in tables:
            add(("table", {**t, "keep": True}))
        add(("h2", "6.2 Charts"))
        descriptive_figures = [f for f in figures if f.get("section", "descriptive") == "descriptive"]
        for fig in descriptive_figures:
            add(("figure", fig))
        if not descriptive_figures:
            add(("p", "No descriptive charts were included in the report. Add charts in the Visualization studio and switch on Add to report."))

    elif key == "inferential":
        add(("h2", "7.1 Summary of decisions"))
        if hyps:
            n_tests = hyps[0].get("tests_in_report") or 0
            add(("p", "'Reject H0' means the data show a difference or link that is unlikely to be chance. "
                      "'Fail to reject H0' means the data do not show one strongly enough."))
            add(("table", {"keep": True, "title": "Decision for each hypothesis", "columns": ["Hypothesis", "Test", "Statistic", "p-value", "Adjusted p (Holm)", "Decision"],
                           "rows": [[h["code"], h.get("test") or "Not tested", h.get("statistic_text") or "-", h.get("p_value"), h.get("p_adjusted"),
                                     h["decision_text"]] for h in hyps],
                           "note": (f"Adjusted p-values allow for the {n_tests} tests run together in this report (Holm method)."
                                    if n_tests > 1 else None)}))
            for h in hyps:
                if h.get("interpretation"):
                    add(("bullet", f"{h['code']}: {h['interpretation']}"))
        else:
            add(("p", "No hypotheses have been defined."))
        if not tests:
            add(("p", "No statistical test has been run."))
        known = {h["code"] for h in hyps}
        for i, h in enumerate(hyps, 1):
            own = [t for t in tests if t.get("hypothesis_code") == h["code"]]
            add(("h2", f"7.{i + 1} {h['code']}: {own[-1]['method_name'] if own else 'Not yet tested'}"))
            add(("p", f"Null hypothesis (H0): {h['null_text']} Alternative hypothesis (H1): {h['alternative_text']}"))
            if own:
                _inferential_blocks(own, add, f"7.{i + 1}.")
                adjusted = h.get("p_adjusted")
                add(("p", f"Decision: {h['decision_text']}" + (f" (adjusted p = {fmt(adjusted)} after the Holm adjustment)" if adjusted is not None else "") + "."))
            else:
                add(("p", "No test has been run for this hypothesis yet."))
        other = [t for t in tests if t.get("hypothesis_code") not in known]
        if other:
            add(("h2", f"7.{len(hyps) + 2} Other analyses"))
            _inferential_blocks(other, add, f"7.{len(hyps) + 2}.")
        figure_no = len(hyps) + (2 if other else 1) + 1
        inferential_figures = [f for f in figures if f.get("section") == "inferential"]
        if inferential_figures:
            add(("h2", f"7.{figure_no} Charts"))
            for fig in inferential_figures:
                add(("figure", fig))

    elif key == "review_findings":
        reviews = content.get("review_findings") or []
        if not reviews:
            add(("p", "No statistical findings have been recorded yet, so there is nothing to review."))
        for review in reviews:
            _review_lines(add, review["lines"], review.get("chart"))

    elif key == "discussion":
        rqs = content.get("research_questions") or []
        if not rqs:
            add(("p", "No research questions have been defined."))
        for rq in rqs:
            add(("h2", f"{rq['code']}: {rq['text']}"))
            add(("p", f"Status: {rq['status_text']}."))
            for f in rq.get("findings", []):
                add(("p", f"What the evidence shows: {f.get('interpretation') or f['statement']}"))
            if not rq.get("findings"):
                add(("p", "No finding has been recorded for this question yet."))
        if rqs and _paragraphs(details.get("background")):
            add(("p", "Earlier studies on this topic are summarised in Section 2 and listed in Section 12."))

    elif key == "limitations":
        items = content.get("limitations") or []
        if not items:
            add(("p", "No limitations have been recorded."))
        for item in items:
            status = "" if item.get("status") == "active" else f" [{item.get('status')}: {item.get('status_note') or 'no note'}]"
            add(("bullet", f"{item['description']}{status}"))

    elif key == "conclusions":
        items = content.get("conclusions") or []
        if not items:
            add(("p", "No conclusions have been recorded."))
        for i, item in enumerate(items, 1):
            add(("h2", f"Conclusion {i}"))
            _layout_lines(add, item["text"])
            basis = ", ".join(f"Finding {n}" for n in item.get("finding_numbers", []))
            if basis:
                add(("note", f"Based on: {basis}."))
        general = content.get("general_conclusion")
        if general:
            add(("h2", "General conclusion"))
            add(("p", general["synthesis"]))
            if general["table"]["rows"]:
                add(("table", {"keep": True, "title": "Hypothesis summary", "columns": general["table"]["columns"],
                               "rows": general["table"]["rows"]}))
            add(("p", general["strength"]))
            add(("h2", "Future research"))
            for item in general["future"]:
                add(("bullet", item))
        for block in content.get("plain_summaries") or []:
            add(("h2", f"Plain-language summary: {block['code']}"))
            for line in block["lines"]:
                add(("p", line))

    elif key == "recommendations":
        items = content.get("recommendations") or []
        if not items:
            add(("p", "No recommendations have been recorded."))
        for i, item in enumerate(items, 1):
            add(("h2", f"Recommendation {i}"))
            _layout_lines(add, item["text"])
            basis = ", ".join(f"Finding {n}" for n in item.get("finding_numbers", []))
            if basis:
                add(("note", f"Based on: {basis}."))

    elif key == "references":
        add(("h2", "12.1 Sources cited"))
        entries = details.get("references") or []
        if entries:
            for entry in entries:
                add(("bullet", entry))
        else:
            add(("note", "No sources have been entered yet. Add the references under Report details on the Report page, in APA 7th edition format."))
        add(("h2", "12.2 Statistical methods used"))
        for m in (content.get("methodology") or {}).get("methods") or []:
            add(("bullet", m["name"]))

    elif key == "appendices":
        a = content.get("appendix") or {}
        add(("h2", "Appendix A. Conditions checked for each test"))
        checks = [[item["method_name"], c["name"], c.get("test") or "", c.get("statistic"), c.get("p_value"),
                   "Yes" if c.get("satisfied") is True else "No" if c.get("satisfied") is False else "Judge", c.get("note") or ""]
                  for item in tests for c in item.get("assumptions") or []]
        if checks:
            add(("table", {"keep": True, "title": "Conditions checked for each test", "columns": ["Test", "Condition", "Check used", "Statistic", "p", "Met?", "Note"],
                           "rows": checks}))
        else:
            add(("p", "No conditions were checked, because no test was run."))
        add(("h2", "Appendix B. Data dictionary"))
        add(("p", "The data dictionary lists every variable, what it means, its type, and its range of values."))
        if a.get("data_dictionary"):
            add(("table", {"keep": True, "title": "Data dictionary", "columns": ["Variable", "Label", "Type", "Level", "Unique", "Missing", "Min", "Max", "Mean", "Median", "SD"],
                           "rows": a["data_dictionary"]}))
        else:
            add(("p", "No dataset has been uploaded."))
        add(("h2", "Appendix C. Analysis settings"))
        add(("p", "Each analysis run is listed with its date, the test used, the data version and the settings, so the result can be repeated."))
        if a.get("analysis_config"):
            add(("table", {"keep": True, "title": "Analysis runs", "columns": ["Run", "Date (UTC)", "Method", "Dataset version", "Variables", "Parameters / filters"],
                           "rows": a["analysis_config"]}))
        else:
            add(("p", "No analysis has been run."))
        add(("h2", "Appendix D. Glossary"))
        add(("table", {"keep": True, "title": "Terms used in this report", "columns": ["Term", "What it means"], "rows": [list(item) for item in GLOSSARY]}))
        add(("note", INTEGRITY_NOTE))
    return blocks


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
_PDF_REPLACE = {"η": "eta", "ω": "omega", "ε": "epsilon", "χ": "chi", "α": "alpha", "β": "beta", "ρ": "rho",
                "τ": "tau", "≤": "<=", "≥": ">=", "−": "-", "–": "-", "—": "-", "‘": "'", "’": "'",
                "“": '"', "”": '"', "…": "..."}


def _pdf_text(value: Any) -> str:
    text = fmt(value) if not isinstance(value, str) else value
    for src, dst in _PDF_REPLACE.items():
        text = text.replace(src, dst)
    text = text.encode("cp1252", "replace").decode("cp1252")
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return text.replace("\u00b2", "<super>2</super>")          # drawn as a real superscript in every PDF viewer


def render_pdf(content: dict, images: dict[str, bytes] | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    images = images or {}
    base = getSampleStyleSheet()
    ink, muted, rule, accent = colors.black, colors.HexColor("#444444"), colors.HexColor("#8c8c8c"), colors.black
    styles = {
        "title": ParagraphStyle("title", parent=base["Title"], fontSize=24, leading=30, textColor=ink, alignment=TA_CENTER, spaceAfter=18),
        "cover": ParagraphStyle("cover", parent=base["Normal"], fontSize=12, leading=18, alignment=TA_CENTER, textColor=muted),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontSize=16, leading=20, textColor=ink, spaceBefore=14, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontSize=12, leading=15, textColor=accent, spaceBefore=10, spaceAfter=4),
        "p": ParagraphStyle("p", parent=base["Normal"], fontSize=10, leading=14, spaceAfter=6),
        "bullet": ParagraphStyle("bullet", parent=base["Normal"], fontSize=10, leading=14, leftIndent=14, bulletIndent=4, spaceAfter=3),
        "note": ParagraphStyle("note", parent=base["Normal"], fontSize=9, leading=12, textColor=muted, leftIndent=8, spaceAfter=5),
        "caption": ParagraphStyle("caption", parent=base["Normal"], fontSize=9, leading=12, textColor=ink, spaceBefore=6, spaceAfter=3, fontName="Helvetica-Bold"),
    }
    doc_width = A4[0] - 4 * cm
    counters = {"table": 0, "figure": 0}

    def table_flowable(t: dict):
        counters["table"] += 1
        cols = t["columns"]
        n_cols = max(len(cols), 1)
        size = 8.5 if n_cols <= 6 else 7.5 if n_cols <= 9 else 6.5
        cell = ParagraphStyle("cell", parent=base["Normal"], fontSize=size, leading=size + 2.5)
        head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold", textColor=colors.black)
        data = [[Paragraph(_pdf_text(c), head) for c in cols]]
        for row in t.get("rows") or []:
            cells = list(row) + [None] * (n_cols - len(row))
            data.append([Paragraph(_pdf_text(v), cell) for v in cells[:n_cols]])
        # first column and long text columns get more room
        lengths = [max([len(str(c))] + [len(fmt(r[i])) if i < len(r) else 0 for r in (t.get("rows") or [])[:40]]) for i, c in enumerate(cols)]
        weights = [min(max(length + 2, 9), 34) for length in lengths]     # floor keeps short numeric columns readable
        widths = [doc_width * w / sum(weights) for w in weights]
        table = Table(data, colWidths=widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, ink), ("LINEBELOW", (0, 1), (-1, -1), 0.25, rule), ("BOX", (0, 0), (-1, -1), 0.5, rule),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        parts = [Paragraph(_pdf_text(f"Table {counters['table']}. {t.get('title') or ''}"), styles["caption"]), table]
        if t.get("note"):
            parts.append(Paragraph(_pdf_text(f"Note. {t['note']}"), styles["note"]))
        parts.append(Spacer(1, 6))
        return parts

    story: list = [Spacer(1, 5 * cm), Paragraph(_pdf_text(content.get("title") or "Research Report"), styles["title"]), Spacer(1, 1 * cm)]
    for label, value in _cover_rows(content):
        story.append(Paragraph(f"<b>{_pdf_text(label)}:</b> {_pdf_text(value)}", styles["cover"]))
    story += [Spacer(1, 2 * cm), Paragraph("Research data analysis report", styles["cover"]), PageBreak()]

    for key, title in _sections(content):
        story.append(Paragraph(_pdf_text(title), styles["h1"]))
        for kind, payload in _text_blocks(content, key):
            if kind == "table":
                parts = table_flowable(payload)
                story.extend([KeepTogether(parts)] if len(payload.get("rows") or []) <= 12 else parts)
            elif kind == "figure":
                png = images.get(payload.get("image_key"))
                if not png:
                    continue
                counters["figure"] += 1
                image = Image(io.BytesIO(png))
                scale = min(doc_width / image.imageWidth, (11 * cm) / image.imageHeight)
                image.drawWidth, image.drawHeight = image.imageWidth * scale, image.imageHeight * scale
                parts = [image, Paragraph(_pdf_text(f"Figure {counters['figure']}. {payload.get('title') or ''}"), styles["caption"])]
                if payload.get("note"):
                    parts.append(Paragraph(_pdf_text(f"Note. {payload['note']}"), styles["note"]))
                story.append(KeepTogether(parts))
            elif kind == "bullet":
                story.append(Paragraph(_pdf_text(payload), styles["bullet"], bulletText="•"))
            else:
                story.append(Paragraph(_pdf_text(payload), styles[kind]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawString(2 * cm, 1.2 * cm, _pdf_text(content.get("title") or "")[:90])
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title=content.get("title") or "Research Report", author=(content.get("project") or {}).get("researcher") or "")
    doc.build(story, onLaterPages=footer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def render_docx(content: dict, images: dict[str, bytes] | None = None) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    images = images or {}
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    counters = {"table": 0, "figure": 0}

    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(content.get("title") or "Research Report")
    run.bold, run.font.size = True, Pt(24)
    doc.add_paragraph()
    for label, value in _cover_rows(content):
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.add_run(f"{label}: ").bold = True
        line.add_run(value)
    doc.add_page_break()

    def add_table(t: dict) -> None:
        counters["table"] += 1
        caption = doc.add_paragraph()
        cap = caption.add_run(f"Table {counters['table']}. {t.get('title') or ''}")
        cap.bold, cap.font.size = True, Pt(10)
        cols = t["columns"]
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = Pt(9) if len(cols) <= 7 else Pt(7.5)
        for i, col in enumerate(cols):
            cell = table.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(col))
            r.bold, r.font.size = True, size
        for row in t.get("rows") or []:
            cells = table.add_row().cells
            for i in range(len(cols)):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(fmt(row[i]) if i < len(row) else "")
                r.font.size = size
        if t.get("note"):
            note = doc.add_paragraph()
            n = note.add_run(f"Note. {t['note']}")
            n.italic, n.font.size = True, Pt(9)
        doc.add_paragraph()

    for key, heading in _sections(content):
        doc.add_heading(heading, level=1)
        for kind, payload in _text_blocks(content, key):
            if kind == "h2":
                doc.add_heading(str(payload), level=2)
            elif kind == "table":
                add_table(payload)
            elif kind == "figure":
                png = images.get(payload.get("image_key"))
                if not png:
                    continue
                counters["figure"] += 1
                doc.add_picture(io.BytesIO(png), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                c = cap.add_run(f"Figure {counters['figure']}. {payload.get('title') or ''}")
                c.bold, c.font.size = True, Pt(10)
                if payload.get("note"):
                    n = doc.add_paragraph().add_run(f"Note. {payload['note']}")
                    n.italic, n.font.size = True, Pt(9)
            elif kind == "bullet":
                doc.add_paragraph(str(payload), style="List Bullet")
            elif kind == "note":
                n = doc.add_paragraph().add_run(str(payload))
                n.italic, n.font.size = True, Pt(9.5)
            else:
                doc.add_paragraph(str(payload))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# XLSX / CSV
# --------------------------------------------------------------------------- #
def results_rows(content: dict) -> tuple[list[str], list[list[Any]]]:
    """One row per inferential result - the flat table behind CSV and the 'Results' sheet."""
    header = ["Research question", "Hypothesis", "Method", "Variables", "N", "Statistic", "Value", "df", "p-value", "Significant",
              "Effect size", "Effect size value", "Decision on H0", "Dataset version", "Run date (UTC)"]
    rows = []
    for item in content.get("inferential") or []:
        effect = next((e for e in item.get("effect_sizes") or [] if e.get("value") is not None), {})
        stat = item.get("statistic") or {}
        rows.append([item.get("rq_code") or "", item.get("hypothesis_code") or "", item["method_name"], item["variables_text"], item.get("n"),
                     stat.get("name"), stat.get("value"), item.get("df"), item.get("p_value"), item.get("significant"),
                     effect.get("name"), effect.get("value"),
                     {"reject_null": "Reject H0", "fail_to_reject_null": "Fail to reject H0"}.get(item.get("decision"), ""),
                     item.get("dataset_version"), item.get("run_at")])
    return header, rows


def render_csv(content: dict) -> bytes:
    header, rows = results_rows(content)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows([["" if v is None else v for v in row] for row in rows])
    return ("﻿" + buffer.getvalue()).encode("utf-8")        # BOM so Excel opens UTF-8 correctly


def render_xlsx(content: dict, data_rows: tuple[list[str], list[list[Any]]] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="14213D")

    def write(ws, row: int, values: list[Any], header: bool = False) -> None:
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=None if value is None else (value if isinstance(value, (int, float, bool)) else str(value)))
            if header:
                cell.font, cell.fill = head_font, head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    def autosize(ws) -> None:
        for idx, column in enumerate(ws.columns, 1):
            width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 60)

    def sheet_of_tables(name: str, tables: list[dict]) -> None:
        ws = wb.create_sheet(name)
        row = 1
        for t in tables:
            ws.cell(row=row, column=1, value=t.get("title")).font = Font(bold=True, size=12)
            write(ws, row + 1, t["columns"], header=True)
            for i, data in enumerate(t.get("rows") or []):
                write(ws, row + 2 + i, data)
            row += len(t.get("rows") or []) + 3
            if t.get("note"):
                ws.cell(row=row - 1, column=1, value=f"Note. {t['note']}").font = Font(italic=True)
                row += 1
        autosize(ws)

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = content.get("title") or "Research Report"
    ws["A1"].font = Font(bold=True, size=14)
    for i, (label, value) in enumerate(_cover_rows(content), 3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    d = content.get("dataset") or {}
    base = len(_cover_rows(content)) + 4
    for i, (label, value) in enumerate([("Dataset", d.get("name")), ("Records analysed", d.get("rows")), ("Variables", d.get("columns")),
                                        ("Dataset version", d.get("version_number")),
                                        ("Data quality score (%)", (content.get("data_quality") or {}).get("score"))]):
        ws.cell(row=base + i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=base + i, column=2, value=value)
    ws.cell(row=base + 7, column=1, value=INTEGRITY_NOTE).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 60

    header, rows = results_rows(content)
    sheet_of_tables("Results", [{"title": "Inferential results", "columns": header, "rows": rows}])
    blocks = {key: [p for kind, p in _text_blocks(content, key, raw_tables=True) if kind == "table"] for key in ("cleaning", "descriptive", "appendices")}
    if [t for t in blocks["appendices"] if t.get("title") == "Data dictionary"]:
        sheet_of_tables("Data dictionary", [t for t in blocks["appendices"] if t.get("title") == "Data dictionary"])
    if [t for t in blocks["appendices"] if t.get("title") == "Assumption checks"]:
        sheet_of_tables("Assumption checks", [t for t in blocks["appendices"] if t.get("title") == "Assumption checks"])
    if blocks["cleaning"]:
        sheet_of_tables("Cleaning log", blocks["cleaning"])
    if blocks["descriptive"]:
        sheet_of_tables("Descriptive", blocks["descriptive"])
    detail = []
    for item in content.get("inferential") or []:
        for t in item.get("tables") or []:
            detail.append({**t, "title": f"{item['method_name']} - {t.get('title')}"})
    if detail:
        sheet_of_tables("Test output", detail)
    narrative = []
    for key, label in (("findings", "Findings"), ("conclusions", "Conclusions"), ("recommendations", "Recommendations"), ("limitations", "Limitations")):
        items = content.get(key) or []
        if key == "findings":
            narrative.append({"title": label, "columns": ["#", "Title", "Finding", "Statistical evidence", "Strength of evidence", "Related to"],
                              "rows": [[i, f["title"], f["statement"], f.get("statistical_evidence"), f.get("confidence_level"),
                                        ", ".join(x for x in (f.get("rq_code"), f.get("hypothesis_code")) if x)] for i, f in enumerate(items, 1)]})
        elif key == "limitations":
            narrative.append({"title": label, "columns": ["#", "Category", "Limitation", "Status"],
                              "rows": [[i, x.get("category"), x["description"], x.get("status")] for i, x in enumerate(items, 1)]})
        else:
            narrative.append({"title": label, "columns": ["#", "Text", "Based on findings"],
                              "rows": [[i, x["text"], ", ".join(str(n) for n in x.get("finding_numbers", []))] for i, x in enumerate(items, 1)]})
    sheet_of_tables("Findings", narrative)
    if data_rows:
        ws = wb.create_sheet("Cleaned data")
        write(ws, 1, data_rows[0], header=True)
        for i, row in enumerate(data_rows[1], 2):
            write(ws, i, row)
        ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


RENDERERS = {"pdf": ("application/pdf", "pdf"), "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
             "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"), "csv": ("text/csv", "csv"),
             "zip": ("application/zip", "zip")}


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None).isoformat() + "Z"
