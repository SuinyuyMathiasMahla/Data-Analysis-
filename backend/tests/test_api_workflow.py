"""End-to-end API test of the complete research workflow on the sample dataset:

register -> project -> research questions -> hypotheses -> upload -> data quality
-> cleaning (audited) -> EDA -> descriptive -> method recommendation -> tests
-> AI explanation -> findings -> conclusions -> recommendations -> limitations
-> charts -> PDF/DOCX/XLSX/CSV report -> history -> reproduction.
"""
import io

from tests.conftest import auth_headers, make_project, upload_sample


def _post(client, url, headers, json=None, expected=(200, 201, 202)):
    response = client.post(url, json=json or {}, headers=headers)
    assert response.status_code in expected, f"{url}: {response.status_code} {response.text[:400]}"
    return response.json()


def _clean(client, headers, dataset_id, operation, params, reason):
    return _post(client, f"/api/datasets/{dataset_id}/clean", headers, {"operation": operation, "params": params, "reason": reason})


def test_complete_research_workflow(client):
    h = auth_headers(client, "researcher@example.org", "Grace Hopper")

    # 1-3. project, research questions, hypotheses -------------------------------------------------
    project = make_project(client, h)
    pid = project["id"]
    rq1 = _post(client, f"/api/projects/{pid}/research-questions", h, {"text": "Is there a relationship between daily social media use and GPA?"})
    rq2 = _post(client, f"/api/projects/{pid}/research-questions", h, {"text": "Does GPA differ between levels of social media use?"})
    rq3 = _post(client, f"/api/projects/{pid}/research-questions", h, {"text": "Which factors predict GPA?"})
    assert [rq1["code"], rq2["code"], rq3["code"]] == ["RQ1", "RQ2", "RQ3"]
    h1 = _post(client, f"/api/projects/{pid}/hypotheses", h, {
        "research_question_id": rq1["id"], "null_text": "There is no significant relationship between social media use and GPA.",
        "alternative_text": "There is a significant relationship between social media use and GPA."})
    assert h1["code"] == "H1" and h1["decision"] == "pending"

    # 4. upload: original preserved, working copy created -----------------------------------------
    dataset = upload_sample(client, h, pid)
    did = dataset["id"]
    kinds = [(v["version_number"], v["kind"]) for v in dataset["versions"]]
    assert kinds == [(0, "original"), (1, "working")]
    assert dataset["row_count"] == 243 and dataset["column_count"] == 20
    types = {v["name"]: v["data_type"] for v in dataset["variables"]}
    assert types["gpa"] == "numeric" and types["gender"] == "categorical" and types["student_id"] == "identifier"
    original_checksum = dataset["versions"][0]["checksum_sha256"]
    page = client.get(f"/api/datasets/{did}/preview?limit=5", headers=h).json()
    assert len(page["rows"]) == 5 and page["total_rows"] == 243          # the browser gets a page, not the file

    # 5. data quality --------------------------------------------------------------------------------
    quality = client.get(f"/api/datasets/{did}/quality-report", headers=h).json()
    codes = {i["code"] for i in quality["issues"]}
    assert {"missing_values", "duplicate_rows", "impossible_value", "invalid_numeric", "inconsistent_category", "outliers"} <= codes
    assert quality["summary"]["duplicate_rows"] == 3 and 0 < quality["score"] < 100
    first_score = quality["score"]

    # 6. cleaning: every step is a new version with a recorded reason ---------------------------------
    missing_reason = client.post(f"/api/datasets/{did}/clean", json={"operation": "remove_duplicates", "params": {}, "reason": ""}, headers=h)
    assert missing_reason.status_code == 422
    step = _clean(client, h, did, "remove_duplicates", {}, "Exact duplicate questionnaire submissions")
    assert step["operation"]["rows_before"] == 243 and step["operation"]["rows_after"] == 240
    step = _clean(client, h, did, "remove_invalid", {"column": "age", "min": 15, "max": 80, "action": "set_missing"}, "Impossible ages are data-entry errors")
    assert step["operation"]["cells_changed"] == 2
    assert {"old": "-5", "new": None} == {k: step["operation"]["details"]["changes"][0][k] for k in ("old", "new")}
    _clean(client, h, did, "remove_invalid", {"column": "attendance_pct", "min": 0, "max": 100, "action": "set_missing"}, "Attendance cannot exceed 100%")
    _clean(client, h, did, "convert_type", {"column": "study_hours", "to": "numeric"}, "Text entered in a numeric field")
    _clean(client, h, did, "standardize_categories", {"column": "gender", "case": "title"}, "Inconsistent capitalisation of gender")
    _clean(client, h, did, "standardize_categories", {"column": "department", "case": "title", "mapping": {"Bussiness": "Business"}}, "Spelling of department")
    _clean(client, h, did, "derive", {"new_column": "stress_score", "operation": "mean", "columns": [f"stress_q{i}" for i in range(1, 6)]},
           "Composite stress scale (mean of five items)")
    noop = client.post(f"/api/datasets/{did}/clean", json={"operation": "remove_duplicates", "params": {}, "reason": "Check again for duplicates"}, headers=h)
    assert noop.status_code == 422                                           # nothing to change -> no empty version
    history = client.get(f"/api/datasets/{did}/cleaning-history", headers=h).json()
    assert len(history["operations"]) == 7 and all(op["reason"] for op in history["operations"])
    detail = client.get(f"/api/datasets/{did}", headers=h).json()
    assert detail["current_version_number"] == 8 and detail["row_count"] == 240 and detail["column_count"] == 21
    assert detail["versions"][0]["checksum_sha256"] == original_checksum     # the original is untouched
    assert detail["quality_score"] > first_score
    original_file = client.get(f"/api/datasets/{did}/download?original=true", headers=h)
    assert original_file.status_code == 200 and b"-5" in original_file.content and b"Bussiness" in original_file.content

    # 7. data dictionary -------------------------------------------------------------------------------
    variables = {v["name"]: v for v in client.get(f"/api/datasets/{did}/dictionary", headers=h).json()["variables"]}
    assert "stress_score" in variables and variables["age"]["n_missing"] == 2
    edited = client.put(f"/api/variables/{variables['gpa']['id']}", headers=h,
                        json={"label": "Grade point average", "role": "dependent", "valid_min": 0, "valid_max": 4})
    assert edited.status_code == 200

    # 8-9. EDA and descriptive statistics ---------------------------------------------------------------
    eda = _post(client, f"/api/projects/{pid}/analysis", h, {"analysis_type": "eda", "dataset_id": did})
    assert eda["status"] == "COMPLETED" and eda["results"][0]["result"]["dimensions"]["rows"] == 240
    desc = _post(client, f"/api/projects/{pid}/analysis", h, {"analysis_type": "descriptive", "dataset_id": did,
                                                               "variables": {"variables": ["gpa", "study_hours", "gender"], "group_by": "gender"}})
    gpa = desc["results"][0]["result"]["numeric"]["gpa"]
    assert desc["status"] == "COMPLETED" and gpa["n"] == 240 and 2.0 < gpa["mean"] < 3.5

    # 10. method recommendation (rule engine; no LLM configured in tests) ------------------------------
    rec = _post(client, f"/api/projects/{pid}/ai/recommend", h, {"dataset_id": did, "research_question_id": rq1["id"],
                                                                 "outcome": "gpa", "predictors": ["social_media_hours"]})
    assert rec["recommendation"]["primary"]["method"] in ("pearson", "spearman")
    assert rec["recommendation"]["source"] == "rule_engine" and rec["explanation"]["source"] == "rule_engine"
    assert rec["recommendation"]["why"] and rec["recommendation"]["assumption_checks"]

    # 11. inferential analyses linked to research questions and hypotheses -------------------------------
    corr = _post(client, f"/api/projects/{pid}/analysis", h, {
        "analysis_type": "inferential", "dataset_id": did, "method": "pearson", "research_question_id": rq1["id"], "hypothesis_id": h1["id"],
        "variables": {"x": "social_media_hours", "y": "gpa"}})
    assert corr["status"] == "COMPLETED", corr
    r = corr["results"][0]
    assert r["statistic_name"] == "r" and -0.7 < r["statistic_value"] < -0.2 and r["p_value"] < 0.001 and r["is_significant"] is True
    assert r["result"]["source"] == "statistical_engine" and r["result"]["software"]["scipy"]
    assert corr["config"]["schema"]["gpa"]["type"] == "numeric" and corr["version_number"] == 8      # reproducibility record
    hypotheses = client.get(f"/api/projects/{pid}/hypotheses", headers=h).json()["hypotheses"]
    assert hypotheses[0]["decision"] == "reject_null"

    added = _post(client, f"/api/analysis/{corr['id']}/test", h, {"method": "spearman", "variables": {"x": "social_media_hours", "y": "gpa"}})
    assert added["status"] == "COMPLETED" and len(added["results"]) == 2

    anova = _post(client, f"/api/projects/{pid}/analysis", h, {
        "analysis_type": "inferential", "dataset_id": did, "method": "one_way_anova", "research_question_id": rq2["id"],
        "variables": {"outcome": "gpa", "group": "usage_level"}})
    assert anova["status"] == "COMPLETED" and anova["results"][0]["result"]["post_hoc"]

    regression = _post(client, f"/api/projects/{pid}/analysis", h, {
        "analysis_type": "inferential", "dataset_id": did, "method": "multiple_linear_regression", "research_question_id": rq3["id"],
        "variables": {"outcome": "gpa", "predictors": ["study_hours", "social_media_hours", "attendance_pct"]},
        "filters": [{"column": "age", "op": "not_missing"}]})
    assert regression["status"] == "COMPLETED" and 0.1 < regression["results"][0]["result"]["model_fit"]["r_squared"] < 0.8
    assert regression["results"][0]["result"]["records_after_filters"] == 238

    failed = _post(client, f"/api/projects/{pid}/analysis", h, {"analysis_type": "inferential", "dataset_id": did, "method": "independent_t_test",
                                                                 "variables": {"outcome": "gpa", "group": "department"}})
    assert failed["status"] == "FAILED" and "exactly 2 groups" in failed["error_message"] and failed["results"] == []

    not_implemented = _post(client, f"/api/projects/{pid}/analysis", h, {"analysis_type": "inferential", "dataset_id": did,
                                                                          "method": "survival_analysis", "variables": {}})
    assert not_implemented["status"] == "FAILED" and "not implemented" in not_implemented["error_message"]

    # 12. AI-assisted interpretation: grounded in the engine result, labelled as rule engine --------------
    explanation = _post(client, f"/api/projects/{pid}/ai/explain", h, {"analysis_result_id": r["id"]})
    assert explanation["source"] == "rule_engine" and r["result"]["apa"] in explanation["content"]
    chat = _post(client, f"/api/projects/{pid}/ai/chat", h, {"message": "What is a p-value?"})
    assert chat["conversation_id"] and "probability" in chat["content"]

    # 13. findings with an evidence link ------------------------------------------------------------------
    finding = _post(client, f"/api/projects/{pid}/findings/from-result", h, {"analysis_result_id": r["id"]})
    assert finding["research_question_id"] == rq1["id"] and finding["hypothesis_id"] == h1["id"]
    assert finding["analysis_result_id"] == r["id"] and "r(" in finding["statistical_evidence"] and "H1" in finding["interpretation"]
    _post(client, f"/api/projects/{pid}/findings/from-result", h, {"analysis_result_id": anova["results"][0]["id"]})
    edited = client.put(f"/api/findings/{finding['id']}", json={"statement": finding["statement"] + " (reviewed by the researcher)"}, headers=h).json()
    assert edited["is_edited"] is True
    evaluation = {q["code"]: q for q in client.get(f"/api/projects/{pid}/research-questions", headers=h).json()["research_questions"]}
    assert evaluation["RQ1"]["status"] == "answered" and evaluation["RQ3"]["status"] == "analysed"

    # 14. conclusions and recommendations are built from findings only ------------------------------------
    conclusions = _post(client, f"/api/projects/{pid}/conclusions/generate", h)["conclusions"]
    # one conclusion per research question: RQ1 and RQ2 are answered; RQ3 has an analysis but no finding, and says so
    answered = [c for c in conclusions if c["finding_ids"]]
    assert len(conclusions) == 3 and len(answered) == 2
    assert [c for c in conclusions if not c["finding_ids"]][0]["text"].startswith("CONCLUSION 3 — Research Question RQ3:")
    assert "no finding has been recorded" in [c for c in conclusions if not c["finding_ids"]][0]["text"]
    recommendations = _post(client, f"/api/projects/{pid}/recommendations/generate", h)["recommendations"]
    assert len(recommendations) == 2 and all(len(x["finding_ids"]) == 1 for x in recommendations)

    # 15. limitations: detected, never silently removed -----------------------------------------------------
    limitations = _post(client, f"/api/projects/{pid}/limitations/detect", h)["limitations"]
    categories = {item["category"] for item in limitations}
    assert {"design", "sampling_bias", "self_report", "generalizability"} <= categories
    auto = limitations[0]
    assert client.delete(f"/api/limitations/{auto['id']}", headers=h).status_code == 400
    assert client.put(f"/api/limitations/{auto['id']}/status", json={"status": "not_applicable"}, headers=h).status_code == 400
    assert client.put(f"/api/limitations/{auto['id']}/status", json={"status": "acknowledged", "status_note": "Discussed in chapter 5"}, headers=h).status_code == 200

    # 16. charts -------------------------------------------------------------------------------------------
    chart = _post(client, f"/api/projects/{pid}/visualizations", h, {"chart_type": "regression", "dataset_id": did, "x": "social_media_hours",
                                                                     "y": "gpa", "include_in_report": True})
    image = client.get(chart["image_url"], headers=h)
    assert image.status_code == 200 and image.content[:8] == b"\x89PNG\r\n\x1a\n"
    regenerated = client.put(f"/api/visualizations/{chart['id']}", json={"title": "GPA and social media use"}, headers=h).json()
    assert regenerated["title"] == "GPA and social media use"
    diag = _post(client, f"/api/analysis/{regression['id']}/diagnostic-plot", h)
    assert diag["chart_type"] == "regression_diagnostics"
    bad_chart = client.post(f"/api/projects/{pid}/visualizations", json={"chart_type": "histogram", "dataset_id": did, "x": "no_such_column"}, headers=h)
    assert bad_chart.status_code == 422

    # 17. reports in all four formats -----------------------------------------------------------------------
    signatures = {"pdf": b"%PDF", "docx": b"PK\x03\x04", "xlsx": b"PK\x03\x04", "csv": b"\xef\xbb\xbf"}
    report_id = None
    for fmt, signature in signatures.items():
        payload = {"format": fmt, "dataset_id": did, "include_data": True}
        if report_id:
            payload["report_id"] = report_id
        report = _post(client, f"/api/projects/{pid}/reports", h, payload)
        report_id = report["id"]
        download = client.get(report["latest_version"]["download_url"], headers=h)
        assert download.status_code == 200 and download.content.startswith(signature), fmt
        assert len(download.content) > 500
    report = client.get(f"/api/reports/{report_id}", headers=h).json()
    assert [v["version_number"] for v in report["versions"]] == [1, 2, 3, 4]
    snapshot = client.get(f"/api/reports/{report_id}/versions/1/content", headers=h).json()
    assert snapshot["project"]["researcher"] == "Grace Hopper" and len(snapshot["inferential"]) == 4
    assert snapshot["hypotheses"][0]["decision_text"] == "Reject H0" and len(snapshot["data_quality"]["cleaning"]) == 7
    # 18. client package: refused until the gate passes, then a zip ---------------------------------------
    refused = client.post(f"/api/projects/{pid}/client-package", json={"dataset_id": did}, headers=h)
    assert refused.status_code == 422 and refused.json()["error"]["code"] == "readiness_failed"
    assert "missing_references" in {i["code"] for i in refused.json()["error"]["details"]["issues"]}
    package = client.post(f"/api/projects/{pid}/client-package", headers=h,
                          json={"dataset_id": did, "references": ["Bull, F. C. et al. (2020). WHO guidelines on physical activity. BJSM."]})
    assert package.status_code == 200, package.text
    assert package.headers["content-type"] == "application/zip"
    import zipfile
    assert {"summary.pdf", "report.pdf", "report.docx", "workbook.xlsx", "cleaned_dataset.csv", "README.md"} <= set(zipfile.ZipFile(io.BytesIO(package.content)).namelist())
    names = set(zipfile.ZipFile(io.BytesIO(package.content)).namelist())
    every_deliverable = {"summary.pdf", "summary.docx", "report.pdf", "report.docx", "cleaned_dataset.csv", "data_dictionary.csv",
                         "workbook.xlsx", "analysis_log.json", "reproduce_analysis.py", "dashboard.html", "bi_export.csv",
                         "handover_deck.pptx", "README.md"}
    assert every_deliverable <= names, every_deliverable - names
    csv_text = client.get(f"/api/reports/{report_id}/download?version=4", headers=h).content.decode("utf-8-sig")
    assert "Pearson correlation" in csv_text and "Reject H0" in csv_text

    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(client.get(f"/api/reports/{report_id}/download?version=3", headers=h).content))
    assert {"Summary", "Results", "Findings", "Cleaned data"} <= set(workbook.sheetnames)
    assert workbook["Cleaned data"].max_row == 241

    # 18. history and reproducibility -------------------------------------------------------------------------
    actions = [e["action"] for e in client.get(f"/api/projects/{pid}/history?limit=500", headers=h).json()["history"]]
    for expected in ("project.created", "dataset.uploaded", "dataset.cleaned", "analysis.completed", "visualization.created",
                     "ai.explain", "finding.drafted", "report.generated"):
        assert expected in actions, expected
    copy = _post(client, f"/api/analysis/{corr['id']}/reproduce", h)
    assert copy["status"] == "COMPLETED" and copy["reproduced_from_id"] == corr["id"]
    check = client.get(f"/api/analysis/{copy['id']}", headers=h).json()["reproduction_check"]
    assert check["identical"] is True and len(check["results"]) == 2
    # the dataset has moved on, but the reproduction still used the stored version
    _clean(client, h, did, "drop_missing", {"columns": ["gender"]}, "Gender is required for group comparisons")
    later = _post(client, f"/api/analysis/{corr['id']}/reproduce", h)
    assert later["version_number"] == 8 and client.get(f"/api/analysis/{later['id']}", headers=h).json()["reproduction_check"]["identical"] is True

    # dashboard and progress ---------------------------------------------------------------------------------
    dashboard = client.get("/api/dashboard", headers=h).json()
    assert dashboard["counts"]["projects"] == 1 and dashboard["counts"]["reports"] == 1 and dashboard["counts"]["datasets"] == 1
    progress = client.get(f"/api/projects/{pid}", headers=h).json()["progress"]
    assert progress["percent"] == 100 and all(s["done"] for s in progress["steps"])


def test_manual_entry_json_upload_and_power(client):
    h = auth_headers(client, "manual@example.org")
    pid = make_project(client, h, academic_level="undergraduate")["id"]
    manual = _post(client, f"/api/projects/{pid}/datasets/manual", h, {
        "name": "Pilot", "columns": ["group", "score"], "rows": [["A", 10], ["A", 12], ["A", 11], ["B", 15], ["B", 17], ["B", 16]]})
    assert manual["source_type"] == "manual" and manual["row_count"] == 6
    tiny = _post(client, f"/api/projects/{pid}/analysis", h, {"analysis_type": "inferential", "dataset_id": manual["id"], "method": "pearson",
                                                              "variables": {"x": "score", "y": "group"}})
    assert tiny["status"] == "FAILED" and "Insufficient evidence to perform this analysis reliably" in tiny["error_message"]
    records = b'[{"x": 1, "y": 2.1}, {"x": 2, "y": 3.9}, {"x": 3, "y": 6.2}, {"x": 4, "y": 8.1}, {"x": 5, "y": 9.8}]'
    uploaded = client.post(f"/api/projects/{pid}/datasets", headers=h, files={"file": ("pairs.json", records, "application/json")})
    assert uploaded.status_code == 201 and uploaded.json()["row_count"] == 5
    plan = _post(client, "/api/power", h, {"design": "t_test_independent", "mode": "required_n", "effect_size": 0.5})
    assert plan["required_n"] == 64 and plan["planning"] == "prospective"
    assert client.post("/api/power", json={"design": "t_test_independent", "effect_size": 0}, headers=h).status_code == 422
    methods = client.get("/api/methods", headers=h).json()["methods"]
    assert any(m["key"] == "survival_analysis" and m["implemented"] is False for m in methods)


def test_llm_provider_is_used_and_checked_when_configured(client):
    from app.ai.providers import AIProvider
    from app.ai.service import AIService
    from app.services import ai_runtime

    class Fake(AIProvider):
        name = "fake"

        def complete(self, system, messages, max_tokens=1200):
            return "The correlation is very strong, r = 0.9731, so social media use causes lower grades."

    ai_runtime.set_ai_service(AIService(Fake(model="fake-1")))
    h = auth_headers(client, "llm@example.org")
    pid = make_project(client, h)["id"]
    did = upload_sample(client, h, pid)["id"]
    run = _post(client, f"/api/projects/{pid}/analysis", h, {"analysis_type": "inferential", "dataset_id": did, "method": "spearman",
                                                             "variables": {"x": "social_media_hours", "y": "gpa"}})
    reply = _post(client, f"/api/projects/{pid}/ai/explain", h, {"analysis_result_id": run["results"][0]["id"]})
    assert reply["source"] == "llm" and reply["provider"] == "fake"
    assert reply["grounding"]["grounded"] is False and "0.9731" in reply["notice"]       # the invented number is flagged to the user
    assert client.get("/api/ai/status", headers=h).json()["mode"] == "llm"
