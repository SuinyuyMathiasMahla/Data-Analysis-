"""Profiling, data quality, cleaning (audit trail), evidence and reports."""
import io
import json

import pandas as pd
import pytest

from app.analysis import cleaning, evidence, profiling, quality, visualization
from app.reports import renderers
from app.reports.content import assemble_report_content
from app.statistics import recommender
from app.statistics.registry import run_method
from app.statistics.utils import AnalysisError
from tests.conftest import SAMPLE_CSV


def sample():
    raw = profiling.read_raw(SAMPLE_CSV.read_bytes(), "csv")
    return raw, profiling.infer_schema(raw)


def cleaned():
    raw, schema = sample()
    for operation, params, reason in [
        ("remove_duplicates", {}, "Duplicate submissions"),
        ("remove_invalid", {"column": "age", "min": 15, "max": 80}, "Impossible ages"),
        ("convert_type", {"column": "study_hours", "to": "numeric"}, "Text in a numeric field"),
        ("standardize_categories", {"column": "gender", "case": "title"}, "Capitalisation"),
        ("standardize_categories", {"column": "department", "case": "title", "mapping": {"Bussiness": "Business"}}, "Spelling"),
    ]:
        outcome = cleaning.apply_operation(raw, schema, operation, params, reason)
        raw, schema = outcome.raw, outcome.schema
    return raw, schema


# --------------------------------------------------------------------------- #
# import and type inference
# --------------------------------------------------------------------------- #
def test_type_inference_on_the_sample_dataset():
    raw, schema = sample()
    assert raw.shape == (243, 20)
    assert schema["student_id"]["type"] == "identifier"
    assert schema["gpa"] == {"type": "numeric", "measurement_level": "ratio"}
    assert schema["gender"] == {"type": "categorical", "measurement_level": "nominal"}
    assert schema["stress_q1"] == {"type": "integer", "measurement_level": "ordinal"}      # 1-5 rating
    assert schema["study_hours"]["type"] == "numeric"                                       # one stray 'abc' does not change the type


def test_working_copy_round_trip_is_lossless():
    raw, _ = sample()
    again = profiling.raw_from_csv_bytes(profiling.raw_to_csv_bytes(raw))
    assert again.shape == raw.shape
    assert (again.fillna("<missing>") == raw.fillna("<missing>")).all().all()


def test_readers_for_json_excel_and_manual_entry():
    records = [{"x": 1, "y": "a"}, {"x": 2, "y": None}, {"x": 3, "y": "b"}]
    from_json = profiling.read_raw(json.dumps(records).encode(), "json")
    assert from_json.shape == (3, 2) and from_json["y"].isna().sum() == 1
    buffer = io.BytesIO()
    pd.DataFrame(records).to_excel(buffer, index=False)
    from_excel = profiling.read_raw(buffer.getvalue(), "xlsx")
    assert from_excel["x"].tolist() == ["1", "2", "3"]
    manual = profiling.frame_from_manual(["group", "score"], [["A", 10], ["B", "N/A"]])
    assert manual["score"].tolist()[0] == "10" and manual["score"].isna().sum() == 1     # 'N/A' is read as missing
    with pytest.raises(AnalysisError):
        profiling.read_raw(b"not json", "json")
    with pytest.raises(AnalysisError):
        profiling.frame_from_manual(["a", "b"], [[1]])


def test_duplicate_and_blank_column_names_are_made_unique():
    raw = profiling.read_raw(b"a,a,,b\n1,2,3,4\n", "csv")
    assert list(raw.columns) == ["a", "a.1", "column_3", "b"]      # pandas numbers the repeat; the blank gets a name


# --------------------------------------------------------------------------- #
# data quality
# --------------------------------------------------------------------------- #
def test_quality_engine_finds_the_injected_problems():
    raw, schema = sample()
    report = quality.assess_quality(raw, schema)
    found = {(i["code"], i["column"]): i for i in report["issues"]}
    assert report["summary"]["duplicate_rows"] == 3
    assert found[("impossible_value", "age")]["count"] == 2
    assert found[("impossible_value", "attendance_pct")]["count"] == 1
    assert found[("invalid_numeric", "study_hours")]["examples"][0]["value"] == "abc"
    assert found[("inconsistent_category", "gender")]["count"] == 3
    assert ("similar_categories", "department") in found                 # 'Bussiness' vs 'Business'
    assert found[("missing_values", "sleep_hours")]["count"] == 6
    assert ("outliers", "social_media_hours") in found
    assert ("duplicate_identifier", "student_id") in found
    assert 90 < report["score"] < 100
    assert all(i["suggestion"] is None or i["suggestion"]["operation"] in cleaning.OPERATIONS for i in report["issues"])


def test_outliers_do_not_lower_the_score():
    clean = pd.DataFrame({"x": [str(v) for v in range(1, 41)]})
    with_outlier = clean.copy()
    with_outlier.loc[0, "x"] = "5000"
    schema = {"x": {"type": "integer", "measurement_level": "ratio"}}
    assert quality.assess_quality(clean, schema)["score"] == 100.0
    report = quality.assess_quality(with_outlier, schema)
    assert report["score"] == 100.0 and any(i["code"] == "outliers" for i in report["issues"])


def test_quality_structural_and_range_checks():
    raw = pd.DataFrame({"empty": [None] * 12, "const": ["1"] * 12, "score": [str(v) for v in [5, 6, 7, 5, 6, 7, 5, 6, 7, 5, 6, 99]],
                        "when": ["2024-01-0%d" % d for d in range(1, 10)] + ["2024-01-10", "2024-01-11", "not a date"]}, dtype=object)
    schema = profiling.infer_schema(raw)
    schema["when"] = {"type": "datetime", "measurement_level": "interval"}
    report = quality.assess_quality(raw, schema, {"score": (0, 10)})
    codes = {i["code"] for i in report["issues"]}
    assert {"empty_column", "constant_column", "impossible_value", "invalid_date"} <= codes
    assert report["dimensions"]["structure"]["points_lost"] > 0


# --------------------------------------------------------------------------- #
# cleaning and its audit trail
# --------------------------------------------------------------------------- #
def test_cleaning_never_mutates_its_input_and_logs_every_change():
    raw, schema = sample()
    before = raw.copy(deep=True)
    outcome = cleaning.apply_operation(raw, schema, "remove_invalid", {"column": "age", "min": 0, "max": 120}, "Invalid age value")
    assert raw.equals(before)                                            # the input version is untouched
    assert outcome.cells_changed == 2
    assert outcome.details["changes"][0] == {"row": 18, "column": "age", "old": "-5", "new": None}
    assert outcome.raw.loc[17, "age"] is None


def test_a_reason_is_mandatory():
    raw, schema = sample()
    for reason in ("", "  ", "ok"):
        with pytest.raises(AnalysisError):
            cleaning.apply_operation(raw, schema, "remove_duplicates", {}, reason)


def test_remove_duplicates_and_drop_missing_report_the_rows():
    raw, schema = sample()
    outcome = cleaning.apply_operation(raw, schema, "remove_duplicates", {}, "Duplicate submissions")
    assert (outcome.rows_before, outcome.rows_after) == (243, 240) and outcome.details["rows_removed"] == [241, 242, 243]
    dropped = cleaning.apply_operation(outcome.raw, schema, "drop_missing", {"columns": ["gender"]}, "Gender is required")
    assert dropped.rows_after == 238


def test_imputation_strategies():
    raw, schema = cleaned()
    median = cleaning.apply_operation(raw, schema, "impute", {"column": "sleep_hours", "strategy": "median"}, "Few values missing")
    assert median.cells_changed == 6 and median.raw["sleep_hours"].isna().sum() == 0
    expected = pd.to_numeric(raw["sleep_hours"]).median()
    assert float(median.details["fill_value"]) == pytest.approx(expected)
    mode = cleaning.apply_operation(raw, schema, "impute", {"column": "gender", "strategy": "mode"}, "Few values missing")
    assert mode.raw["gender"].isna().sum() == 0
    with pytest.raises(AnalysisError):
        cleaning.apply_operation(raw, schema, "impute", {"column": "gender", "strategy": "mean"}, "Not possible for text")


def test_standardise_recode_and_replace():
    raw, schema = sample()
    out = cleaning.apply_operation(raw, schema, "standardize_categories", {"column": "department", "case": "title", "mapping": {"Bussiness": "Business"}}, "Spelling")
    assert sorted(out.raw["department"].dropna().unique()) == ["Business", "Computer Science", "Education", "Engineering"]
    recoded = cleaning.apply_operation(raw, schema, "recode", {"column": "passed", "mapping": {"Yes": 1, "No": 0}, "new_column": "passed_num"}, "Numeric coding")
    assert set(recoded.raw["passed_num"].unique()) == {"0", "1"} and "passed" in recoded.raw.columns
    assert recoded.schema["passed_num"]["type"] == "integer"
    replaced = cleaning.apply_operation(raw, schema, "replace_values", {"column": "age", "mapping": {"199": None}}, "Sentinel code")
    assert replaced.cells_changed == 1


def test_numeric_transformations():
    raw, schema = cleaned()
    z = cleaning.apply_operation(raw, schema, "standardize", {"column": "gpa", "new_column": "gpa_z"}, "Standardise")
    values = pd.to_numeric(z.raw["gpa_z"])
    assert values.mean() == pytest.approx(0, abs=1e-6) and values.std(ddof=1) == pytest.approx(1, abs=1e-6)
    norm = cleaning.apply_operation(raw, schema, "normalize", {"column": "gpa", "new_column": "gpa_01"}, "Scale")
    scaled = pd.to_numeric(norm.raw["gpa_01"])
    assert scaled.min() == 0 and scaled.max() == 1
    with pytest.raises(AnalysisError):                                   # log of a non-positive number is refused, not silently NaN
        cleaning.apply_operation(raw.assign(neg="-1"), {**schema, "neg": {"type": "numeric", "measurement_level": "ratio"}},
                                 "transform", {"column": "neg", "function": "log"}, "Try a log")
    wins = cleaning.apply_operation(raw, schema, "winsorize", {"column": "social_media_hours", "lower_pct": 1, "upper_pct": 99}, "Cap one implausible value")
    assert pd.to_numeric(wins.raw["social_media_hours"]).max() < 23.5 and "Winsorization" in wins.details["note"]
    assert all(c["old"] != c["new"] for c in wins.details["changes"])    # only real changes are logged
    scale = cleaning.apply_operation(raw, schema, "derive", {"new_column": "stress", "operation": "mean", "columns": ["stress_q1", "stress_q2"]}, "Scale score")
    first = (float(raw.loc[0, "stress_q1"]) + float(raw.loc[0, "stress_q2"])) / 2
    assert float(scale.raw.loc[0, "stress"]) == pytest.approx(first)


def test_unknown_operations_and_columns_are_rejected():
    raw, schema = sample()
    with pytest.raises(AnalysisError):
        cleaning.apply_operation(raw, schema, "run_python", {"code": "import os"}, "Not allowed")
    with pytest.raises(AnalysisError):
        cleaning.apply_operation(raw, schema, "impute", {"column": "nope"}, "Unknown column")


# --------------------------------------------------------------------------- #
# recommender
# --------------------------------------------------------------------------- #
def test_recommendations_follow_variable_types_and_assumptions():
    raw, schema = cleaned()
    df = profiling.typed_frame(raw, schema)
    cases = [
        (dict(outcome="gpa", predictors=["gender"]), {"independent_t_test", "mann_whitney"}),
        (dict(outcome="gpa", predictors=["department"]), {"one_way_anova", "kruskal_wallis"}),
        (dict(outcome="passed", predictors=["gender"]), {"chi_square", "fisher_exact"}),
        (dict(outcome="passed", predictors=["study_hours"]), {"logistic_regression"}),
        (dict(outcome="stress_q1", predictors=["sleep_hours"]), {"spearman"}),
        (dict(outcome="gpa", predictors=["study_hours", "attendance_pct"]), {"multiple_linear_regression"}),
        (dict(outcome="library_visits", predictors=["study_hours", "gender"], outcome_kind="count"), {"poisson_regression", "negative_binomial_regression"}),
        (dict(paired=True, measures=["pre_test", "post_test"]), {"paired_t_test", "wilcoxon"}),
        (dict(measures=["stress_q1", "stress_q2", "stress_q3"]), {"friedman"}),
    ]
    for kwargs, allowed in cases:
        rec = recommender.recommend(df, schema, **kwargs)
        assert rec["primary"]["method"] in allowed, (kwargs, rec["primary"])
        assert rec["why"] and rec["source"] == "rule_engine"
        result = run_method(df, schema, rec["config"]["method"], rec["config"]["variables"]) if rec["primary"]["implemented"] and rec["primary"]["method"] != "negative_binomial_regression" else None
        assert result is None or result["n"] > 0                          # the recommended configuration actually runs


def test_recommender_is_honest_about_missing_methods_and_level():
    raw, schema = cleaned()
    df = profiling.typed_frame(raw, schema)
    survival = recommender.recommend(df, schema, outcome="gpa", predictors=["gender"], outcome_kind="time_to_event")
    assert survival["primary"]["implemented"] is False and any("not implemented" in w for w in survival["warnings"])
    undergrad = recommender.recommend(df, schema, outcome="gpa", predictors=["study_hours"], academic_level="undergraduate")
    phd = recommender.recommend(df, schema, outcome="gpa", predictors=["study_hours"], academic_level="phd")
    assert undergrad["primary"]["method"] == phd["primary"]["method"]       # the level never changes the method...
    assert len(phd["level_notes"]) > len(undergrad["level_notes"])          # ...only the depth of the guidance
    with pytest.raises(AnalysisError):
        recommender.recommend(df, schema, outcome="student_id", predictors=["gpa"])


# --------------------------------------------------------------------------- #
# evidence trail, charts and reports
# --------------------------------------------------------------------------- #
def _result():
    raw, schema = cleaned()
    df = profiling.typed_frame(raw, schema)
    return df, schema, run_method(df, schema, "pearson", {"x": "social_media_hours", "y": "gpa"})


def test_findings_quote_the_engine_and_never_claim_causation():
    _, _, result = _result()
    hypothesis = {"id": 1, "code": "H1", "direction": "non_directional"}
    draft = evidence.draft_finding(result, {"id": 1, "code": "RQ1", "text": "?"}, hypothesis)
    assert draft["statement"] == result["summary"] and result["apa"] in draft["statistical_evidence"]
    assert "H1: the null hypothesis is rejected" in draft["interpretation"]
    assert "does not show that one variable causes the other" in draft["statement"]
    assert draft["confidence_level"] in ("low", "moderate", "high") and draft["limitations_note"]
    non_significant = {**result, "p_value": 0.4, "decision": "fail_to_reject_null"}
    assert "not evidence that H0 is true" in evidence.hypothesis_verdict(non_significant, hypothesis)


def test_conclusions_and_recommendations_are_tied_to_findings():
    findings = [{"id": 7, "title": "A", "statement": "S1.", "research_question_id": 1, "confidence_level": "low", "significant": True},
                {"id": 8, "title": "B", "statement": "S2.", "research_question_id": None, "confidence_level": "high", "significant": False},
                {"id": 9, "title": "C", "statement": "Manual.", "research_question_id": 1, "confidence_level": "moderate", "significant": None}]
    conclusions = evidence.draft_conclusions(findings, [{"id": 1, "code": "RQ1", "text": "Is X related to Y?"}])
    assert conclusions[0]["finding_ids"] == [7, 9] and "tentative" in conclusions[0]["text"]
    assert conclusions[1]["finding_ids"] == [8]
    recommendations = evidence.draft_recommendations(findings, "Cross-sectional survey")
    assert [r["finding_ids"] for r in recommendations] == [[7], [8]]           # nothing is invented for the manual finding
    assert recommendations[0]["text"].startswith("RECOMMENDATION 1: Audience:") and "Expected measurable outcome" in recommendations[0]["text"]
    assert "do not change practice" in recommendations[1]["text"]


def test_limitations_are_detected_from_design_data_and_results():
    _, _, result = _result()
    project = {"study_design": "Cross-sectional survey", "sampling_technique": "Convenience sampling", "data_collection_method": "Questionnaire"}
    found = evidence.detect_limitations(project, {"rows": 60, "missing_pct": 8.0}, [result] * 5, [{"operation": "winsorize"}])
    categories = {item["category"] for item in found}
    assert {"sample_size", "missing_data", "design", "sampling_bias", "self_report", "distribution", "multiple_testing",
            "data_handling", "generalizability"} <= categories


def test_every_chart_type_renders_a_png():
    raw, schema = cleaned()
    df = profiling.typed_frame(raw, schema)
    configs = [dict(chart_type="bar", x="department"), dict(chart_type="bar", x="department", y="gpa"), dict(chart_type="pie", x="usage_level"),
               dict(chart_type="histogram", x="gpa"), dict(chart_type="box", x="department", y="gpa"), dict(chart_type="violin", x="gender", y="gpa"),
               dict(chart_type="scatter", x="study_hours", y="gpa", hue="gender"), dict(chart_type="regression", x="social_media_hours", y="gpa"),
               dict(chart_type="line", x="year_of_study", y="gpa"), dict(chart_type="heatmap", x="department", y="usage_level"),
               dict(chart_type="correlation_matrix", columns=["gpa", "study_hours", "sleep_hours"]),
               dict(chart_type="group_comparison", x="usage_level", y="gpa", category_order=["Low", "Moderate", "High"]), dict(chart_type="qq", x="gpa")]
    assert {c["chart_type"] for c in configs} | {"time_series"} == set(visualization.CHART_TYPES)
    for config in configs:
        png, meta = visualization.render_chart(df, config, {"gpa": "Grade point average"})
        assert png[:8] == b"\x89PNG\r\n\x1a\n" and meta["title"] and meta["n"] > 0, config
    dated = pd.DataFrame({"day": pd.date_range("2024-01-01", periods=30), "value": range(30)})
    assert visualization.render_chart(dated, {"chart_type": "time_series", "x": "day", "y": "value"})[0][:4] == b"\x89PNG"
    with pytest.raises(AnalysisError):
        visualization.render_chart(df, {"chart_type": "histogram", "x": "missing_column"})
    with pytest.raises(AnalysisError):
        visualization.render_chart(df, {"chart_type": "time_series", "x": "age", "y": "gpa"})


def test_report_renders_in_all_formats_from_one_snapshot():
    df, schema, result = _result()
    png, meta = visualization.render_chart(df, {"chart_type": "regression", "x": "social_media_hours", "y": "gpa"})
    finding = {**evidence.draft_finding(result), "id": 1, "research_question_id": 1, "hypothesis_id": 1, "analysis_result_id": 5, "is_edited": False}
    content = assemble_report_content(
        title="Test report", researcher="Ada Lovelace", project={"title": "T", "academic_level": "masters", "study_design": "Cross-sectional survey"},
        dataset={"name": "Survey", "source_type": "csv", "original_filename": "s.csv", "version_number": 5, "rows": 240, "columns": 20, "original_checksum": "ab" * 32},
        variables=[{"name": "gpa", "data_type": "numeric", "measurement_level": "ratio", "n_missing": 0, "n_unique": 100, "stats": {"mean": 2.6}}],
        quality={"score": 99.2, "version_number": 5, "summary": {"missing_cells": 3}, "issues": []},
        cleaning=[{"operation": "remove_duplicates", "reason": "Duplicates", "rows_before": 243, "rows_after": 240, "cells_changed": 0}],
        runs=[{"id": 3, "name": "Pearson", "status": "COMPLETED", "research_question_id": 1, "version_number": 5, "completed_at": "2026-01-01T10:00:00Z",
               "results": [{"id": 5, "result_type": "test", "result": result, "test": {"hypothesis_id": 1, "params": {"alpha": 0.05}}}]},
              {"id": 4, "name": "Failed run", "status": "FAILED", "results": []}],
        research_questions=[{"id": 1, "code": "RQ1", "text": "Is social media use related to GPA?"}],
        hypotheses=[{"id": 1, "code": "H1", "null_text": "No relationship.", "alternative_text": "A relationship.", "alpha": 0.05}],
        findings=[finding], conclusions=[{"text": "Conclusion.", "finding_ids": [1], "source": "engine"}],
        recommendations=[{"text": "Recommendation.", "finding_ids": [1], "source": "manual"}],
        limitations=[{"category": "design", "description": "Cross-sectional.", "status": "acknowledged", "status_note": "Discussed", "source": "auto"}],
        figures=[{"image_key": "f1", "title": meta["title"], "note": None, "section": "inferential"}])
    assert len(content["inferential"]) == 1 and content["inferential"][0]["rq_code"] == "RQ1"
    assert content["hypotheses"][0]["decision_text"] == "Reject H0" and content["research_questions"][0]["status"] == "answered"
    assert content["conclusions"][0]["finding_numbers"] == [1]
    json.dumps(content)                                                    # the snapshot must be storable as JSON

    pdf = renderers.render_pdf(content, {"f1": png})
    docx = renderers.render_docx(content, {"f1": png})
    xlsx = renderers.render_xlsx(content, (["gpa"], [["2.5"], ["3.1"]]))
    csv_bytes = renderers.render_csv(content)
    assert pdf.startswith(b"%PDF") and docx.startswith(b"PK") and xlsx.startswith(b"PK")
    text = csv_bytes.decode("utf-8-sig")
    assert "Pearson correlation" in text and "Reject H0" in text and "RQ1" in text

    from docx import Document
    from openpyxl import load_workbook
    document = Document(io.BytesIO(docx))
    body = "\n".join(p.text for p in document.paragraphs)
    for heading in ("1. Abstract", "2. Background and literature review", "3. Research questions and hypotheses", "5. Data cleaning summary",
                    "6. Results: descriptive statistics", "7. Results: statistical tests", "7.1 Summary of decisions", "9. Limitations",
                    "10. Conclusions", "12. References", "Appendix A", "Appendix B", "Appendix D"):
        assert heading in body
    assert result["apa"] in body and "Ada Lovelace" in body and len(document.inline_shapes) == 1
    workbook = load_workbook(io.BytesIO(xlsx))
    assert workbook["Results"]["C3"].value == "Pearson correlation" and workbook["Cleaned data"].max_row == 3
    only = renderers.render_docx({**content, "sections": ["inferential"]}, {})
    only_text = "\n".join(p.text for p in Document(io.BytesIO(only)).paragraphs)
    assert "7. Results: statistical tests" in only_text and "2. Background and literature review" not in only_text


def test_primary_hypothesis_is_raw_and_secondary_hypotheses_are_holm_adjusted():
    def test_result(result_id, hypothesis_id, p, variables="x vs y"):
        payload = {"method": "pearson", "method_name": "Pearson correlation", "variables": {"x": "x", "y": "y"}, "p_value": p, "n": 50,
                   "statistic": {"name": "r", "value": 0.3}, "significant": p < 0.05,
                   "decision": "reject_null" if p < 0.05 else "fail_to_reject_null", "summary": "Summary text.", "apa": "r = 0.30"}
        return {"id": result_id, "result_type": "test", "result": payload, "test": {"hypothesis_id": hypothesis_id, "params": {}}}

    run = {"id": 1, "name": "Runs", "status": "COMPLETED", "research_question_id": 1, "version_number": 1,
           "completed_at": "2026-01-01T10:00:00Z", "results": [test_result(1, 1, 0.03), test_result(2, 2, 0.04), test_result(3, 3, 0.2),
                                                             test_result(4, 1, 0.03)]}   # 4 is a re-run of test 1: counted once
    content = assemble_report_content(
        title="T", researcher=None, project={}, dataset=None, variables=[], quality=None, cleaning=[], runs=[run],
        research_questions=[{"id": 1, "code": "RQ1", "text": "Q?"}],
        hypotheses=[{"id": i, "code": f"H{i}", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05} for i in (1, 2, 3)],
        findings=[], conclusions=[], recommendations=[], limitations=[], figures=[])

    assert len(content["inferential"]) == 3
    by_code = {h["code"]: h for h in content["hypotheses"]}
    assert by_code["H1"]["primary"] and by_code["H1"]["p_adjusted"] is None and by_code["H1"]["decision_text"] == "Reject H0"
    assert by_code["H2"]["p_adjusted"] == 0.08 and by_code["H2"]["decision_text"] == "Fail to reject H0"
    assert by_code["H3"]["p_adjusted"] == 0.2 and by_code["H3"]["tests_in_report"] == 2
    assert "Holm-adjusted p = 0.080" in by_code["H2"]["interpretation"]


def test_client_package_gate_blocks_then_builds_a_complete_zip():
    import zipfile
    from app.reports import gate, package

    def test_result(result_id, hypothesis_id, p):
        payload = {"method": "pearson", "method_name": "Pearson correlation", "variables": {"x": "x", "y": "y"}, "p_value": p, "n": 50,
                   "statistic": {"name": "r", "value": 0.3}, "significant": p < 0.05,
                   "decision": "reject_null" if p < 0.05 else "fail_to_reject_null", "summary": "Summary.", "apa": "r = 0.30",
                   "effect_sizes": [{"name": "r", "value": 0.3, "ci_low": 0.1, "ci_high": 0.5}]}
        return {"id": result_id, "result_type": "test", "result": payload, "test": {"hypothesis_id": hypothesis_id, "params": {}}}

    run = {"id": 1, "name": "Runs", "status": "COMPLETED", "version_number": 1, "completed_at": "2026-01-01T10:00:00Z",
           "results": [test_result(1, 1, 0.03), test_result(2, 2, 0.2)]}           # hypothesis 3 has no result
    content = assemble_report_content(
        title="T", researcher=None, project={}, dataset=None, variables=[], quality=None, cleaning=[], runs=[run],
        research_questions=[{"id": 1, "code": "RQ1", "text": "Q?"}],
        hypotheses=[{"id": i, "code": f"H{i}", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05} for i in (1, 2, 3)],
        findings=[], conclusions=[], recommendations=[], limitations=[], figures=[])

    codes = {i["code"] for i in gate.check(content, references=[])}
    assert {"hypothesis_without_result", "missing_references"} <= codes
    blocking = {i["code"] for i in gate.check(content, references=["Author, A. (2020). A title. Journal."])}
    assert blocking == {"hypothesis_without_result"}

    content["hypotheses"][2]["test"] = "Pearson correlation"                    # resolve the missing result
    content["hypotheses"][2]["p_value"] = 0.4
    content["hypotheses"][2]["p_adjusted"] = 0.4
    assert gate.check(content, references=["Author, A. (2020). A title. Journal."]) == []

    data = package.build(content, {}, ["summary", "report", "log"], frame=None, references=["Author, A. (2020). A title. Journal."])
    names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
    assert {"summary.pdf", "report.pdf", "report.docx", "analysis_log.json", "README.md"} <= names
    assert zipfile.ZipFile(io.BytesIO(data)).read("summary.pdf").startswith(b"%PDF")


def test_gate_blocks_a_number_in_the_text_that_no_table_shows():
    from app.reports import gate

    def test_result(result_id, hypothesis_id, p):
        payload = {"method": "pearson", "method_name": "Pearson correlation", "variables": {"x": "x", "y": "y"}, "p_value": p, "n": 50,
                   "statistic": {"name": "r", "value": 0.3}, "significant": p < 0.05,
                   "decision": "reject_null" if p < 0.05 else "fail_to_reject_null", "summary": "Summary.", "apa": f"r = 0.30, p = {p:.3f}"}
        return {"id": result_id, "result_type": "test", "result": payload, "test": {"hypothesis_id": hypothesis_id, "params": {}}}

    run = {"id": 1, "name": "Runs", "status": "COMPLETED", "version_number": 1, "completed_at": "2026-01-01T10:00:00Z",
           "results": [test_result(1, 1, 0.035)]}
    content = assemble_report_content(
        title="T", researcher=None, project={}, dataset=None, variables=[], quality=None, cleaning=[], runs=[run],
        research_questions=[{"id": 1, "code": "RQ1", "text": "Q?"}],
        hypotheses=[{"id": 1, "code": "H1", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05}],
        findings=[], conclusions=[{"text": "Men scored higher, p = 0.420.", "finding_ids": [], "source": "manual"}],
        recommendations=[], limitations=[], figures=[])
    issues = [i for i in gate.check(content, references=["A. (2020). B."]) if i["code"] == "number_mismatch"]
    assert len(issues) == 1 and "0.420" in issues[0]["message"]
    content["conclusions"][0]["text"] = "There was an association, p = 0.035."
    content["executive_summary"] = ["Overall conclusion: There was an association, p = 0.035."]
    assert not [i for i in gate.check(content, references=["A. (2020). B."]) if i["code"] == "number_mismatch"]


def test_one_page_summary_verdicts_and_word_copy():
    import io as _io
    from docx import Document
    from app.reports import package

    def test_result(result_id, hypothesis_id, p):
        payload = {"method": "pearson", "method_name": "Pearson correlation", "variables": {"x": "x", "y": "y"}, "p_value": p, "n": 50,
                   "statistic": {"name": "r", "value": 0.3}, "significant": p < 0.05,
                   "decision": "reject_null" if p < 0.05 else "fail_to_reject_null", "summary": "S.", "apa": "r = 0.30"}
        return {"id": result_id, "result_type": "test", "result": payload, "test": {"hypothesis_id": hypothesis_id, "params": {}}}

    run = {"id": 1, "name": "R", "status": "COMPLETED", "research_question_id": 1, "version_number": 1,
           "completed_at": "2026-01-01T10:00:00Z", "results": [test_result(1, 1, 0.01), test_result(2, 2, 0.4)]}
    content = assemble_report_content(
        title="T", researcher=None, project={}, dataset=None, variables=[], quality=None, cleaning=[], runs=[run],
        research_questions=[{"id": 1, "code": "RQ1", "text": "Q1?"}, {"id": 2, "code": "RQ2", "text": "Q2?"}],
        hypotheses=[{"id": 1, "code": "H1", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05},
                    {"id": 2, "code": "H2", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05}],
        findings=[], conclusions=[], recommendations=[], limitations=[], figures=[])
    content["inferential"][1]["rq_code"] = "RQ2"
    content["inferential"][0]["rq_code"] = "RQ1"
    assert package._verdict(content, "RQ1") == "Yes"
    assert package._verdict(content, "RQ2") == "Inconclusive"           # not rejected: never 'No' automatically
    word = Document(_io.BytesIO(package.summary_docx(content, "executive")))
    text = "\n".join(p.text for p in word.paragraphs) + "\n".join(c.text for t in word.tables for r in t.rows for c in r.cells)
    assert "Prepared for: decision-makers" in text and "Verdict by research question" in text
    assert package.summary_pdf(content, "academic").startswith(b"%PDF")


def test_dashboard_and_bi_export_are_self_contained_and_escape_data():
    import json as _json
    from app.reports import package

    def test_result(result_id, hypothesis_id, p, name):
        payload = {"method": "pearson", "method_name": name, "variables": {"x": "x", "y": "y"}, "p_value": p, "n": 50,
                   "statistic": {"name": "r", "value": 0.3}, "significant": p < 0.05,
                   "decision": "reject_null" if p < 0.05 else "fail_to_reject_null", "summary": "S.", "apa": "r = 0.30",
                   "effect_sizes": [{"name": "r", "value": 0.3, "ci_low": 0.1, "ci_high": 0.5}]}
        return {"id": result_id, "result_type": "test", "result": payload, "test": {"hypothesis_id": hypothesis_id, "params": {}}}

    run = {"id": 1, "name": "R", "status": "COMPLETED", "research_question_id": 1, "version_number": 1,
           "completed_at": "2026-01-01T10:00:00Z", "results": [test_result(1, 1, 0.01, "</script><b>x</b>"), test_result(2, 2, 0.4, "Plain test")]}
    content = assemble_report_content(
        title="T", researcher=None, project={}, dataset=None, variables=[], quality=None, cleaning=[], runs=[run],
        research_questions=[{"id": 1, "code": "RQ1", "text": "Q1?"}],
        hypotheses=[{"id": 1, "code": "H1", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05},
                    {"id": 2, "code": "H2", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05}],
        findings=[], conclusions=[], recommendations=[], limitations=[], figures=[])
    page = package.dashboard_html(content).decode("utf-8")
    assert "</script><b>" not in page                       # a test name cannot close the data script and inject markup
    data = page.split('<script id="data" type="application/json">')[1].split("</script>")[0]
    tests = _json.loads(data.replace("<\/", "</"))["tests"]
    assert {t["decision"] for t in tests} == {"Reject H0", "Fail to reject H0"}
    csv_text = package.bi_export_csv(content).decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("question,hypothesis,test,variables") and len(csv_text.splitlines()) == 3


def test_handover_deck_has_8_to_12_slides_with_notes():
    import io as _io
    from pptx import Presentation
    from app.reports import package

    def test_result(result_id, hypothesis_id, p):
        payload = {"method": "pearson", "method_name": "Pearson correlation", "variables": {"x": "study_time", "y": "gpa"}, "p_value": p,
                   "n": 50, "statistic": {"name": "r", "value": 0.3}, "significant": p < 0.05,
                   "decision": "reject_null" if p < 0.05 else "fail_to_reject_null", "summary": "Study time and GPA are related.",
                   "apa": "r = 0.30", "effect_sizes": [{"name": "r", "value": 0.3, "ci_low": 0.1, "ci_high": 0.5}]}
        return {"id": result_id, "result_type": "test", "result": payload, "test": {"hypothesis_id": hypothesis_id, "params": {}}}

    run = {"id": 1, "name": "R", "status": "COMPLETED", "research_question_id": 1, "version_number": 1,
           "completed_at": "2026-01-01T10:00:00Z", "results": [test_result(1, 1, 0.01)]}
    content = assemble_report_content(
        title="Study time and GPA", researcher=None, project={}, dataset={"name": "Survey", "rows": 200, "columns": 9, "version_number": 1},
        variables=[], quality=None, cleaning=[{"operation": "remove_duplicates", "reason": "Exact copies", "rows_before": 210, "rows_after": 200,
                                               "cells_changed": 0}],
        runs=[run], research_questions=[{"id": 1, "code": "RQ1", "text": "Is study time related to GPA?"}],
        hypotheses=[{"id": 1, "code": "H1", "null_text": "No association.", "alternative_text": "An association.", "alpha": 0.05}],
        findings=[], conclusions=[], recommendations=[{"text": "Offer study skills sessions.", "finding_ids": [], "source": "manual",
                                                       "audience": "Head of department"}],
        limitations=[{"category": "design", "description": "Cross-sectional.", "status": "active", "source": "auto"}], figures=[])
    data = package.deck_pptx(content, {})
    deck = Presentation(_io.BytesIO(data))
    assert 8 <= len(deck.slides) <= 12, len(deck.slides)
    assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in deck.slides)
    text = " ".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    assert "RQ1" in text and "Head of department" in text and "Cross-sectional" in text


def test_reproduce_script_recomputes_a_logged_correlation(tmp_path):
    import subprocess
    import sys as _sys
    import zipfile
    import numpy as np
    import pandas as pd
    from scipy import stats
    from app.reports import package

    rng = np.random.default_rng(3)
    x = rng.normal(size=60)
    y = 0.4 * x + rng.normal(size=60)
    frame = pd.DataFrame({"study_time": x, "gpa": y})
    p = round(float(stats.pearsonr(x, y)[1]), 6)
    payload = {"method": "pearson", "method_name": "Pearson correlation", "variables": {"x": "study_time", "y": "gpa"}, "p_value": p,
               "n": 60, "statistic": {"name": "r", "value": 0.3}, "significant": True, "decision": "reject_null", "summary": "S.", "apa": "r"}
    run = {"id": 1, "name": "R", "status": "COMPLETED", "research_question_id": 1, "version_number": 1, "completed_at": "2026-01-01T10:00:00Z",
           "results": [{"id": 1, "result_type": "test", "result": payload, "test": {"hypothesis_id": 1, "params": {}}}]}
    content = assemble_report_content(
        title="T", researcher=None, project={}, dataset=None, variables=[], quality=None, cleaning=[], runs=[run],
        research_questions=[{"id": 1, "code": "RQ1", "text": "Q?"}],
        hypotheses=[{"id": 1, "code": "H1", "null_text": "n.", "alternative_text": "a.", "alpha": 0.05}],
        findings=[], conclusions=[], recommendations=[], limitations=[], figures=[])
    data = package.build(content, {}, ["log", "data"], frame=frame, references=[], audience="academic")
    folder = tmp_path / "package"
    zipfile.ZipFile(io.BytesIO(data)).extractall(folder)
    result = subprocess.run([_sys.executable, "reproduce_analysis.py"], cwd=folder, capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr
    assert "match" in result.stdout and "DIFFERS" not in result.stdout
